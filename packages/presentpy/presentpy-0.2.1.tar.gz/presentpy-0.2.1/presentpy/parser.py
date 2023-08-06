import mistletoe
import nbformat
import pkg_resources
from mistletoe import block_token
from pptx import Presentation

from presentpy.code import get_config_from_source, get_parsed_lines
from presentpy.slides import add_bullet_slide, add_code_slide, add_title_slide

BLANK_TEMPLATE = pkg_resources.resource_filename("presentpy", "templates/Blank.pptx")


def process_notebook(file):
    presentation = Presentation(BLANK_TEMPLATE)
    with open(file) as r:
        notebook = nbformat.read(r, as_version=4)

        for cell in notebook["cells"]:
            source = cell["source"]
            if not source:
                continue

            if cell["cell_type"] == "markdown":
                process_markdown_cell(source, presentation)
            elif cell["cell_type"] == "code":
                process_code_cell(source, presentation)
    return presentation


def process_code_cell(source, presentation):
    source, cell_config = get_config_from_source(source)
    parsed_lines = get_parsed_lines(source)
    add_code_slide(presentation, parsed_lines, cell_config)
    return source


def process_markdown_cell(source, presentation):
    document = mistletoe.Document(source)
    header = document.children[0]
    if len(document.children) > 1:
        if isinstance(document.children[1], block_token.Heading):
            sub_header = document.children[1]
            add_title_slide(presentation, header.children[0].content, sub_header.children[0].content, header.level)
        elif isinstance(document.children[1], block_token.List):
            bullets = [bullet.children[0].children[0].content for bullet in document.children[1].children]
            add_bullet_slide(
                presentation,
                header.children[0].content,
                bullets,
            )

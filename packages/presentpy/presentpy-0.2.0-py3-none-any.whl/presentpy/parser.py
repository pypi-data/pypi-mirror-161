import mistletoe
import nbformat
from mistletoe import block_token
from pptx import Presentation

from presentpy.code import get_config_from_source, get_parsed_lines
from presentpy.slides import add_bullet_slide, add_code_slide, add_title_slide


def process_notebook(file):
    presentation = Presentation()
    with open(file) as r:
        notebook = nbformat.read(r, as_version=4)

        for cell in notebook["cells"]:
            source = cell["source"]
            if not source:
                continue

            if cell["cell_type"] == "markdown":
                process_markdown_cell(cell, presentation)
            elif cell["cell_type"] == "code":
                process_code_cell(cell, presentation)
    return presentation


def process_code_cell(source, presentation):
    source_lines = source.split("\n")
    cell_config = get_config_from_source(source_lines)
    source = "\n".join(source_lines[:-1])
    parsed_lines = get_parsed_lines(source)
    add_code_slide(presentation, parsed_lines, cell_config)
    return source


def process_markdown_cell(source, presentation):
    document = mistletoe.Document(source)
    header = document.children[0]
    if len(document.children) > 1:
        if isinstance(document.children[1], block_token.Heading):
            sub_header = document.children[1]
            add_title_slide(presentation, header.children[0].content, sub_header.children[0].content)
        elif isinstance(document.children[1], block_token.List):
            bullets = [bullet.children[0].children[0].content for bullet in document.children[1].children]
            add_bullet_slide(
                presentation,
                header.children[0].content,
                bullets,
            )
